#!/usr/bin/env python3
"""Build the episode's two-page downloadable knowledge map with python-docx."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.55), Cm(1.55)
sec.left_margin, sec.right_margin = Cm(1.8), Cm(1.8)
for name, size in [('Normal', 10), ('Title', 19), ('Subtitle', 11), ('Heading 1', 13), ('Heading 2', 11)]:
    st = doc.styles[name]
    st.font.name = 'Arial Unicode MS'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.1
    for border in list(st.element.xpath('./w:pPr/w:pBdr')):
        border.getparent().remove(border)
    if name.startswith('Heading'):
        st.paragraph_format.space_before = Pt(8)
        st.paragraph_format.keep_with_next = True

def p(text, style=None):
    return doc.add_paragraph(text, style)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9')
        borders.append(e)
    props.append(borders)
    for j, text in enumerate(headers):
        t.rows[0].cells[j].text = text
    rep = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        cells = t.add_row().cells
        for j, text in enumerate(row): cells[j].text = text
    for i, row in enumerate(t.rows):
        no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for j, cell in enumerate(row.cells):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'bottom', 'left', 'right'):
                e = OxmlElement('w:' + side); e.set(qn('w:w'), '75'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            cp.append(margins)
            if i == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E4EDF2'); cp.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(.5)
                para.paragraph_format.space_before = Pt(.5)
                for run in para.runs:
                    run.font.size = Pt(9.5)
                    if i == 0: run.bold = True
    return t

p('第2集 维京船上的寻物单', 'Title')
p('名词与冠词知识地图', 'Subtitle')
p('找宝物时，先说清是哪一个，再说谁的、有几个。忘记规则时，可以从下面的例句查起。')

p('a 和 an 看紧接的读音', 'Heading 1')
p('一个没指定的可数东西，常用 a 或 an。看紧接的声音，不只看第一个字母。')
table(['用 a', '用 an'], [
    ('a shield 一面盾牌', 'an axe 一把斧头'),
    ('a useful tool 一个有用的工具\nuseful 开头是辅音 /j/', 'an hour 一小时\nhour 开头的 h 不发音'),
    ('a small helmet 一顶小头盔\n紧接的是 small', 'an old helmet 一顶旧头盔\n紧接的是 old'),
], [8.7, 8.7])

p('the 让对方认出哪一个', 'Heading 1')
p('听的人能认出是哪一个，就常用 the。不一定要先提过；同看的画面、后面的描述，都能帮助认出来。')
p('Bring an axe. The axe is too heavy!\n拿一把斧头。那把斧头太重了！')
p("Bring me the helmet in Biscuit's mouth.\n把 Biscuit 嘴里那顶头盔拿来。")
p("Please pass the bread on Lily's plate.\n请递来 Lily 盘里那份面包。")

p('这些表达通常不加冠词', 'Heading 1')
table(['表达的意思', '例句'], [
    ('泛指复数或不可数名词', 'Dogs like socks. / We need water.'),
    ('通常的三餐 球类 棋类 学科', 'after lunch / play football / play chess\nlearn English'),
    ('已经有 my 等限定词', 'my helmet'),
], [6.1, 11.3])
p('有指定或修饰时，表达可以改变：a big lunch 是一顿丰盛的午餐；the lunch on the table 指桌上的那顿午餐。')

p('再看一眼容易混淆的地方', 'Heading 1')
p("an axe，但 a useful tool；说某一份面包时可以用 the bread。\n说足球这项运动用 Leo plays football.\ncaptain's 是一位船长的；sailors' 是多位船员的。")

doc.add_page_break()
p('名词告诉我们是什么和有多少', 'Title')
p('先看这个词在句子里的意思，再决定能不能直接数。', 'Subtitle')
p('可以直接数物件，也可以数单位', 'Heading 1')
p('名词给人和东西起名字。Leo 是专有名词；helmet 是普通名词。可数与不可数是另一种分类。')
table(['直接数物件', '按材料或食物的份量说'], [
    ('two shields 两面盾牌\nthree knives 三把小刀\nsocks 袜子', 'bread 面包 / water 水 / wood 木材'),
    ('多个可数东西，通常用复数', 'two pieces of bread 两块面包\na bottle of water 一瓶水'),
], [8.7, 8.7])
p('同一个词，意思变了，可数性也可能变。room 表“空间”时不可数；a room 表“一间房间”。不要用“能不能切开”来判断。')

p('多个东西的词形变化', 'Heading 1')
table(['样式', '例子'], [
    ('一般加 s', 'shield → shields'),
    ('部分词加 es', 'box → boxes'),
    ('辅音字母加 y 结尾', 'city → cities；但 key → keys'),
    ('部分 f 或 fe 结尾变 ves', 'knife → knives；但 roof → roofs'),
    ('部分 o 结尾加 es', 'potato → potatoes；但 photo → photos'),
    ('不规则变化或单复数相同', 'child → children；man → men\nsheep → sheep'),
], [6.1, 11.3])

p('所有格说明是谁的', 'Heading 1')
table(['主人', '表达'], [
    ('一位主人', "the captain's helmet / Leo's shield"),
    ('以 s 结尾的复数主人', "the sailors' boots"),
    ('不以 s 结尾的复数主人', "the children's bags"),
], [6.1, 11.3])
p("my shield 也能说明主人。无生命事物不只用 of：the deck of the ship 和 the ship's deck 都可以。")

doc.core_properties.title = '第2集 名词与冠词知识地图'
doc.core_properties.subject = '罗盘冒险队 维京船上的寻物单'
doc.core_properties.author = '罗盘冒险队'
doc.core_properties.keywords = '名词 冠词 所有格 复数'
target = ROOT / 'downloads' / 'ep02_nouns_articles.docx'
target.parent.mkdir(parents=True, exist_ok=True)
doc.save(target)
check = Document(target)
assert len(check.tables) == 5
assert any("the sailors' boots" in c.text for t in check.tables for r in t.rows for c in r.cells)
print(target)
